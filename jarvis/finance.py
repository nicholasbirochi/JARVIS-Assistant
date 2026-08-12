"""Reads the user's own manually-maintained net-worth spreadsheet
(Patrimônio.xlsx) and produces a local financial summary -- a stand-in
for real bank integration (Meu Pluggy, still pending the user creating
that account themselves) that doesn't need one: this file already lives
on the user's own machine, no credentials, no external API, same
"100% local" principle as everything else in jarvis/. Never transmits
anything -- load_snapshot() just reads a local file with openpyxl.

Real spreadsheet shape (found live 2026-08-11 by reading the actual
file, not guessed -- and cross-checked: every derived total below was
verified to exactly match the sheet's own "Contas" tab before trusting
this parsing): two sheets, "Variáveis" and "Contas".

"Variáveis" -- row 1 is a title, row 2 is blank, row 3 is a section
label, row 4 is the real column header, data starts row 5:
- Columns A-C: one row per income entry (Tipo/Dinheiro/Mês-Ano) --
  "Salário" plus occasional one-offs ("Reembolso", "Venda X"). Rows
  extend into the future with a date pre-filled and Dinheiro empty (a
  planning template, not real data) -- excluded by requiring Dinheiro to
  be numeric, not just present.
- Columns E-G: one row per one-off expense (Tipo/Dinheiro/Quantidade) --
  NOT aligned month-to-month with the income rows; independent list
  sharing the same sheet. The real cost is Dinheiro * Quantidade, not
  Dinheiro alone -- confirmed by cross-checking against Contas'
  "Despesas" total, which only matched once quantity was factored in
  (e.g. "GymPass" 69.9 * 18 real months, not a flat 69.9).
- Columns I-V: a single snapshot row (real data's row 1, i.e. sheet row
  5) holding current per-bank figures across three groups of four bank
  columns (Inter BR, Inter US, Itaú, Mercado Pago): "Dinheiro Investido"
  (principal), "Rendimento/Moeda" (yield), "Dinheiro Total" (both
  summed). Confirmed live: Inter US's Investido/Rendimento are in raw
  USD, but its Total is already BRL-converted (the numbers don't add up
  in USD, but do at a plausible USD/BRL rate) -- so Total, not
  Investido + Rendimento, is the trustworthy BRL figure to sum across
  banks for a net-worth total.

"Contas" -- a tiny 2-row aggregate (header row + one values row):
Valor Bruto (= sum of real income entries), Despesas (= -sum(Dinheiro *
Quantidade) across expense entries), Valor Previsto (= Valor Bruto +
Despesas), Luxos (= Grana Final - Valor Previsto -- the gap between the
bare-bones salary-minus-expenses prediction and what's actually
accumulated, i.e. net investment yield + unaccounted side income),
Grana Final (= sum of every bank's "Dinheiro Total", the real net-worth
figure). These are already-computed formula RESULTS (openpyxl loaded
with data_only=True) -- read directly, not re-derived, since Contas is
what the exact matches above were checked against.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

DEFAULT_PATH = Path(
    "/Users/nicholasbirochi/Library/CloudStorage/OneDrive-FundaçãoSalvadorArena/Extras/Finanças/Patrimônio.xlsx"
)

BANKS = ["Inter BR", "Inter US", "Itaú", "Mercado Pago"]

# Inter US is the one bank column that holds USD, not BRL, for its raw
# Investido/Rendimento figures (confirmed live -- see module docstring).
# Total is already BRL-converted for every bank, including this one.
_USD_BANKS = {"Inter US"}


@dataclass
class IncomeEntry:
    kind: str
    amount: float
    when: datetime


@dataclass
class ExpenseEntry:
    kind: str
    unit_amount: float
    quantity: float

    @property
    def total(self) -> float:
        return self.unit_amount * self.quantity


@dataclass
class BankSnapshot:
    bank: str
    invested: float
    yield_: float
    total_brl: float
    currency: str  # "BRL" or "USD" -- describes invested/yield_, total_brl is always BRL


@dataclass
class FinanceSnapshot:
    income: list[IncomeEntry] = field(default_factory=list)
    expenses: list[ExpenseEntry] = field(default_factory=list)
    banks: list[BankSnapshot] = field(default_factory=list)
    gross_value: float | None = None  # "Valor Bruto"
    total_expenses: float | None = None  # "Despesas" (negative)
    projected_value: float | None = None  # "Valor Previsto"
    luxuries: float | None = None  # "Luxos"
    final_money: float | None = None  # "Grana Final" -- the net-worth total


def load_snapshot(path: Path | str = DEFAULT_PATH) -> FinanceSnapshot:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    snapshot = FinanceSnapshot()
    _parse_variaveis(wb["Variáveis"], snapshot)
    _parse_contas(wb["Contas"], snapshot)
    return snapshot


def _parse_variaveis(ws, snapshot: FinanceSnapshot) -> None:
    for row in ws.iter_rows(min_row=5, values_only=True):
        kind, amount, when = row[0], row[1] if len(row) > 1 else None, row[2] if len(row) > 2 else None
        if kind and isinstance(amount, (int, float)) and isinstance(when, datetime):
            snapshot.income.append(IncomeEntry(kind=kind, amount=float(amount), when=when))

        if len(row) > 6:
            ex_kind, ex_amount, ex_qty = row[4], row[5], row[6]
            if ex_kind and isinstance(ex_amount, (int, float)):
                qty = float(ex_qty) if isinstance(ex_qty, (int, float)) else 1.0
                snapshot.expenses.append(ExpenseEntry(kind=ex_kind, unit_amount=float(ex_amount), quantity=qty))

        if len(row) > 21 and not snapshot.banks and any(v is not None for v in row[8:12]):
            invested, yield_, total = row[8:12], row[13:17], row[18:22]
            for i, bank in enumerate(BANKS):
                snapshot.banks.append(
                    BankSnapshot(
                        bank=bank,
                        invested=float(invested[i] or 0),
                        yield_=float(yield_[i] or 0),
                        total_brl=float(total[i] or 0),
                        currency="USD" if bank in _USD_BANKS else "BRL",
                    )
                )


def _parse_contas(ws, snapshot: FinanceSnapshot) -> None:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return
    values = rows[1]
    snapshot.gross_value = values[0]
    snapshot.total_expenses = values[1]
    snapshot.projected_value = values[3]
    snapshot.luxuries = values[4]
    snapshot.final_money = values[6]


def expenses_by_category(snapshot: FinanceSnapshot, *, top_n: int = 5) -> list[tuple[str, float]]:
    """Highest-cost expense entries (not grouped by repeated kind -- in
    this sheet each row is already its own one-off purchase, so there's
    rarely more than one row per kind to group)."""
    return sorted(((e.kind, e.total) for e in snapshot.expenses), key=lambda pair: pair[1], reverse=True)[:top_n]


def format_brl(value: float) -> str:
    """1234.5 -> "1.234,50" -- Brazilian thousands/decimal separators,
    swapped from Python's US-style `:,.2f` output rather than pulling in
    locale machinery for two separator characters."""
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def summarize_finances(snapshot: FinanceSnapshot) -> str:
    """Short text summary meant for a voice/text reply."""
    lines = []
    if snapshot.final_money is not None:
        lines.append(f"Patrimônio total: R$ {format_brl(snapshot.final_money)}.")
    if snapshot.banks:
        bank_parts = []
        for b in snapshot.banks:
            if b.currency == "USD" and b.invested:
                bank_parts.append(f"{b.bank}: R$ {format_brl(b.total_brl)} (US$ {format_brl(b.invested)} + rendimento)")
            else:
                bank_parts.append(f"{b.bank}: R$ {format_brl(b.total_brl)}")
        lines.append("Por conta: " + "; ".join(bank_parts) + ".")
    if snapshot.gross_value is not None and snapshot.total_expenses is not None:
        lines.append(
            f"Ganhou R$ {format_brl(snapshot.gross_value)} no total registrado, gastou "
            f"R$ {format_brl(-snapshot.total_expenses)}."
        )
    if snapshot.luxuries is not None:
        lines.append(f"Rendimento/ganho além do previsto: R$ {format_brl(snapshot.luxuries)}.")
    top_expenses = expenses_by_category(snapshot, top_n=3)
    if top_expenses:
        formatted = ", ".join(f"{kind} (R$ {format_brl(total)})" for kind, total in top_expenses)
        lines.append(f"Maiores gastos: {formatted}.")
    return " ".join(lines) if lines else "Não consegui ler dados financeiros da planilha."
